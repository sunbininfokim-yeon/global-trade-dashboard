"""Observed public series for the Japan indicators that were still fixture_synth.

Sources (keyless unless noted):
  - MHLW spreadsheets found through the e-Stat data catalog (needs ESTAT_APP_ID; the files themselves
    download without a key): the seasonally adjusted job-openings-to-applicants ratio and the real
    wage growth (nominal cash earnings deflated by CPI excl. imputed rent, 5+ employees);
  - Statistics Bureau of Japan: the monthly CPI files (national and Tokyo ku-area) published as CSV
    on stat.go.jp -- no e-Stat API key needed;
  - Bank of Japan Time-Series Data Search API (stat-search.boj.or.jp): the BOJ's own accounts, the
    call rate, M2, the producer price index, the balance of payments;
  - Ministry of Finance: the daily JGB yield curve (jgbcme_all.csv);
  - FRED: real and nominal GDP (Cabinet Office), foreign reserves;
  - CFTC: the yen futures positioning report.

Same contract as us_public_series.py: every value is a published observation or arithmetic on one,
a series that cannot be read leaves the previous card, nothing is estimated.

What is deliberately not here, and why: PMIs and CDS are licensed data; TOPIX has no free feed (Yahoo has
no index symbol; a TOPIX ETF is a different quantity); the Nikkei VI history is downloadable but
carries Nikkei's "do not copy or circulate" notice; stooq answers with a bot-check page; the JPX
investor-type spreadsheets have no terms page that could be found to confirm reuse; the Rengo
shunto tally page could not be located; the MOF publishes the
intervention detail quarterly, so a month it has not covered yet cannot be told apart from a month
with no intervention.
"""

from __future__ import annotations

import csv
import io
import json
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any, Callable

from . import us_public_series as ups
from .us_public_series import Points, Spec

BOJ_API = "https://www.stat-search.boj.or.jp/api/v1/getDataCode"
MOF_JGB = "https://www.mof.go.jp/english/policy/jgbs/reference/interest_rate/historical/jgbcme_all.csv"
MOF_JGB_MONTH = "https://www.mof.go.jp/english/policy/jgbs/reference/interest_rate/jgbcme.csv"     # the current month, daily
CFTC_YEN = "https://publicreporting.cftc.gov/resource/6dca-aqww.json"
ESTAT_CATALOG = "https://api.e-stat.go.jp/rest/3.0/app/json/getDataCatalog"
ESTAT_CREDIT = "이 서비스는 정부통계 종합창구(e-Stat)의 API 기능을 사용하고 있지만, 서비스의 내용은 국가가 보증한 것이 아닙니다."
STAT_CPI = "https://www.stat.go.jp/data/cpi/{base}/csv/{kind}{base}aa.csv"     # kind: zmi = national, tmi = Tokyo ku-area
UA = "macro-monitor/1.0 (+https://github.com/sunbininfokim-yeon/global-trade-dashboard)"

# display formats the shared module does not have
ups._FORMATS.update({
    "pct0": lambda v: f"{v:.0f}%",
    "pct2": lambda v: f"{v:.2f}%",
    "tn1": lambda v: f"{v:.1f}T",
    "tn1s": lambda v: f"{v:+.1f}T",
    "bn0u": lambda v: f"{v:,.0f}B",
    "bp0": lambda v: f"{v:+.0f}bp" if v else "0bp",
    "k0s": lambda v: f"{v:+,.0f}K",
    "ratio2": lambda v: f"{v:.2f}",
})


def _get(url: str, *, timeout: int = 60, tries: int = 3) -> bytes:
    last: Exception | None = None
    for attempt in range(tries):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": UA}), timeout=timeout) as resp:
                return resp.read()
        except (urllib.error.URLError, TimeoutError) as exc:
            last = exc
            time.sleep(2 * (attempt + 1))
    raise RuntimeError(f"{url.split('?')[0]}: {last}")


# --------------------------------------------------------------------------
# Bank of Japan
# --------------------------------------------------------------------------

def parse_boj_csv(text: str) -> Points:
    """The API's CSV: a status header, then SERIES_CODE,NAME,UNIT,FREQUENCY,CATEGORY,LAST_UPDATE,SURVEY_DATES,VALUES.
    Monthly survey dates are YYYYMM; a value that is empty or not a number is skipped."""
    if "STATUS,200" not in text[:200]:
        raise ValueError(f"BOJ API did not answer OK: {text[:120]!r}")
    i = text.find("SERIES_CODE")
    if i < 0:
        raise ValueError("no data header in the BOJ answer")
    out: Points = []
    for row in csv.DictReader(io.StringIO(text[i:])):
        sd, v = (row.get("SURVEY_DATES") or "").strip(), (row.get("VALUES") or "").strip()
        if len(sd) != 6 or v in ("", "NA", "-"):
            continue
        try:
            out.append((f"{sd[:4]}-{sd[4:6]}-01", float(v)))
        except ValueError:
            continue
    if not out:
        raise ValueError("BOJ series has no observations")
    return sorted(out)


def fetch_boj(db: str, code: str, *, start: str = "201501") -> Points:
    url = f"{BOJ_API}?format=csv&lang=en&db={db}&code={urllib.parse.quote(code, safe='')}&startDate={start}"
    return parse_boj_csv(_get(url).decode("utf-8", errors="replace"))


# --------------------------------------------------------------------------
# Ministry of Finance: JGB yields
# --------------------------------------------------------------------------

def parse_mof_jgb(text: str, tenors: tuple[str, ...] = ("2Y", "10Y", "30Y")) -> dict[str, Points]:
    """Daily par yields by tenor from jgbcme_all.csv (Date,1Y,2Y,...; dates Y/M/D; '-' when a tenor was not quoted)."""
    lines = text.lstrip("﻿").splitlines()
    hdr = next((i for i, ln in enumerate(lines) if ln.startswith("Date,")), None)
    if hdr is None:
        raise ValueError("no 'Date,' header in the MOF yield file")
    cols = [c.strip() for c in lines[hdr].split(",")]
    idx = {t: cols.index(t) for t in tenors}
    out: dict[str, Points] = {t: [] for t in tenors}
    for row in csv.reader(lines[hdr + 1:]):
        if not row or "/" not in row[0]:
            continue
        y, m, d = (int(x) for x in row[0].strip().split("/"))
        iso = f"{y:04d}-{m:02d}-{d:02d}"
        for t, i in idx.items():
            if i < len(row) and row[i].strip() not in ("", "-"):
                out[t].append((iso, float(row[i])))
    for t, pts in out.items():
        if not pts:
            raise ValueError(f"MOF file has no {t} yields")
    return out


def month_last(points: Points) -> Points:
    """The last observation of each month (the current month's is the latest one so far)."""
    last: dict[str, tuple[str, float]] = {}
    for d, v in points:
        last[d[:7]] = (d, v)
    return [last[k] for k in sorted(last)]


def fetch_mof_jgb() -> dict[str, Points]:
    """The full history file is refreshed monthly; the current month's daily file has the days since.
    Both are read and merged by date (the newer file wins where they overlap)."""
    hist = parse_mof_jgb(_get(MOF_JGB, timeout=120).decode("utf-8", errors="replace"))
    try:
        cur = parse_mof_jgb(_get(MOF_JGB_MONTH, timeout=60).decode("utf-8", errors="replace"))
    except (ValueError, RuntimeError):
        cur = {t: [] for t in hist}                  # a month that has no quoted day yet has no data rows
    return {t: sorted({**dict(hist[t]), **dict(cur[t])}.items()) for t in hist}


# --------------------------------------------------------------------------
# CFTC: yen futures, non-commercial net position
# --------------------------------------------------------------------------

def parse_cftc_yen(rows: list[dict[str, Any]]) -> Points:
    """Non-commercial long minus short, contracts, by report date (a Tuesday), ascending."""
    out = []
    for r in rows:
        d = str(r.get("report_date_as_yyyy_mm_dd", ""))[:10]
        try:
            out.append((d, float(r["noncomm_positions_long_all"]) - float(r["noncomm_positions_short_all"])))
        except (KeyError, ValueError):
            continue
    if not out:
        raise ValueError("no CFTC yen rows")
    return sorted(out)


def fetch_cftc_yen(weeks: int = 560) -> Points:
    q = urllib.parse.urlencode({"$where": "cftc_contract_market_code='097741'", "$order": "report_date_as_yyyy_mm_dd DESC", "$limit": weeks})
    return parse_cftc_yen(json.loads(_get(f"{CFTC_YEN}?{q}").decode("utf-8")))


# --------------------------------------------------------------------------
# Statistics Bureau: CPI (index levels by month, one column per item)
# --------------------------------------------------------------------------

def parse_stat_cpi_csv(text: str, items: tuple[str, ...]) -> dict[str, Points]:
    """Monthly index levels for the named items (English header row). The file is one row per month
    (YYYYMM) and one column per item; a name that is missing or appears more than once among the
    columns raises rather than picking one."""
    rows = list(csv.reader(io.StringIO(text)))
    hdr = next((r for r in rows if r and r[0].startswith("Group/Item")), None)
    if hdr is None:
        raise ValueError("no 'Group/Item' header row in the CPI file")
    idx: dict[str, int] = {}
    for name in items:
        hits = [i for i, h in enumerate(hdr) if h.strip() == name]
        if len(hits) != 1:
            raise ValueError(f"CPI column {name!r} found {len(hits)} times")
        idx[name] = hits[0]
    out: dict[str, Points] = {n: [] for n in items}
    for r in rows:
        if r and len(r[0]) == 6 and r[0].isdigit():
            for name, i in idx.items():
                if i < len(r) and r[i].strip():
                    out[name].append((f"{r[0][:4]}-{r[0][4:]}-01", float(r[i])))
    for name, pts in out.items():
        if not pts:
            raise ValueError(f"CPI column {name!r} has no rows")
    return out


def fetch_stat_cpi(kind: str, base: int, items: tuple[str, ...]) -> dict[str, Points]:
    text = _get(STAT_CPI.format(base=base, kind=kind), timeout=90).decode("cp932", errors="replace")
    return parse_stat_cpi_csv(text, items)


def chained_yoy(old_base: Points, new_base: Points, *, switch: str) -> Points:
    """Year-on-year change (%) of the index. Up to the month before `switch` it is taken on the old
    base; from `switch` on it is taken on the new base, as the Bureau does after a re-basing. The
    new-base file starts at `switch`'s previous year, so it can give a YoY from `switch` on."""
    old = {d: v for d, v in old_base}
    new = {d: v for d, v in new_base}
    out: Points = []
    for d in sorted(old):
        prev = f"{int(d[:4]) - 1}-{d[5:]}"
        if d < switch and prev in old:
            out.append((d, (old[d] / old[prev] - 1) * 100))
    for d in sorted(new):
        prev = f"{int(d[:4]) - 1}-{d[5:]}"
        if d >= switch and prev in new:
            out.append((d, (new[d] / new[prev] - 1) * 100))
    return out


# --------------------------------------------------------------------------
# MHLW spreadsheets through the e-Stat data catalog
# --------------------------------------------------------------------------

class MissingKey(RuntimeError):
    """ESTAT_APP_ID is not set: the two e-Stat backed cards are skipped, not failed."""


def recent_months(today, n: int = 5) -> list[str]:
    """YYYYMM for this month and the n-1 before it (a release month whose file is not out yet has no catalog entry)."""
    out, y, m = [], today.year, today.month
    for _ in range(n):
        out.append(f"{y:04d}{m:02d}")
        y, m = (y - 1, 12) if m == 1 else (y, m - 1)
    return out


def estat_find_file(app_id: str, *, search_word: str, title_has: str, months: list[str], stats_code: str | None = None,
                    get: Callable[[str], bytes] | None = None) -> tuple[str, str]:
    """(file url, YYYYMM) of the newest catalog resource whose title contains `title_has`."""
    get = get or (lambda u: _get(u, timeout=120))
    for ym in months:
        params = {"appId": app_id, "searchWord": search_word, "surveyYears": ym, "limit": 20}
        if stats_code:
            params["statsCode"] = stats_code
        d = json.loads(get(f"{ESTAT_CATALOG}?{urllib.parse.urlencode(params)}").decode("utf-8"))["GET_DATA_CATALOG"]
        if str(d["RESULT"]["STATUS"]) not in ("0", "1"):                      # 1 = answered fine, nothing for that month
            raise RuntimeError(f"e-Stat catalog: {d['RESULT'].get('ERROR_MSG')}")
        entries = (d.get("DATA_CATALOG_LIST_INF") or {}).get("DATA_CATALOG_INF") or []
        for entry in ([entries] if isinstance(entries, dict) else entries):
            res = entry["RESOURCES"]["RESOURCE"]
            for r in ([res] if isinstance(res, dict) else res):
                t = r.get("TITLE")
                name = t.get("NAME") if isinstance(t, dict) else t
                if name and title_has in name:
                    return r["URL"], ym
    raise RuntimeError(f"no e-Stat file titled {title_has!r} in {months[-1]}..{months[0]}")


def parse_job_ratio_xlsx(data: bytes) -> Points:
    """MHLW long time-series table 3: one row per year, 12 monthly columns for the actual figures then 12
    for the seasonally adjusted ones (the headline ratio is the adjusted one, all workers incl. part-time)."""
    import unicodedata

    import openpyxl

    ws = openpyxl.load_workbook(io.BytesIO(data), data_only=True)[openpyxl.load_workbook(io.BytesIO(data), read_only=True).sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = next((i for i, r in enumerate(rows) if r and r[0] == "西暦"), None)
    if hdr is None:
        raise ValueError("no '西暦' header row in the job-ratio table")
    blocks, cols = rows[hdr - 1], {}
    for j, label in enumerate(rows[hdr]):
        m = unicodedata.normalize("NFKC", str(label or "")).strip()
        if blocks[j] == "季節調整値" and m.endswith("月") and m[:-1].isdigit() and "～" not in m:
            cols[int(m[:-1])] = j
    if sorted(cols) != list(range(1, 13)):
        raise ValueError(f"seasonally adjusted month columns not found: {sorted(cols)}")
    out: Points = []
    for r in rows[hdr + 1:]:
        y = unicodedata.normalize("NFKC", str(r[0] or ""))
        if y.endswith("年") and y[:-1].isdigit():
            for mo, j in cols.items():
                if isinstance(r[j], (int, float)):
                    out.append((f"{int(y[:-1])}-{mo:02d}-01", float(r[j])))
    if not out:
        raise ValueError("job-ratio table has no observations")
    return sorted(out)


def parse_real_wage_xls(data: bytes) -> Points:
    """MHLW long time-series table 25-1 (real cash earnings, all industries, 5+ employees): a level section
    then a 'year-on-year' section (前年比), one row per year, months in columns 8-19."""
    import xlrd

    return parse_real_wage_sheet(xlrd.open_workbook(file_contents=data).sheet_by_name("TL"))


def parse_real_wage_sheet(sh) -> Points:
    start = next((r for r in range(sh.nrows) if str(sh.cell_value(r, 0)).strip().startswith("前年比")), None)
    if start is None:
        raise ValueError("no year-on-year section in the real-wage table")
    head = next(r for r in range(start, sh.nrows) if str(sh.cell_value(r, 0)).strip() == "年")
    months = {int(float(sh.cell_value(head, c))): c for c in range(sh.ncols)
              if isinstance(sh.cell_value(head, c), float) and 1 <= sh.cell_value(head, c) <= 12}
    if sorted(months) != list(range(1, 13)):
        raise ValueError(f"month columns not found in the year-on-year section: {sorted(months)}")
    out: Points = []
    for r in range(head + 1, sh.nrows):
        y = sh.cell_value(r, 0)
        if not isinstance(y, float):
            continue
        for mo, c in months.items():
            v = sh.cell_value(r, c)
            if isinstance(v, float):
                out.append((f"{int(y)}-{mo:02d}-01", v))
    if not out:
        raise ValueError("real-wage table has no observations")
    return sorted(out)


def fetch_estat_file(kind: str, today=None) -> tuple[Points, dict[str, str]]:
    """(series, where it came from) -- needs ESTAT_APP_ID for the catalog lookup only."""
    import os
    from datetime import datetime, timezone

    app_id = os.environ.get("ESTAT_APP_ID", "").strip()
    if not app_id:
        raise MissingKey("ESTAT_APP_ID is not set")
    months = recent_months(today or datetime.now(timezone.utc).date(), 5)
    if kind == "job_ratio":
        url, ym = estat_find_file(app_id, search_word="一般職業紹介状況", title_has="_3_有効求人倍率", months=months)
        return parse_job_ratio_xlsx(_get(url, timeout=120)), {"file": url.replace(app_id, "***"), "catalog_month": ym}
    if kind == "real_wage":
        url, ym = estat_find_file(app_id, search_word="長期時系列表", title_has="25-1_実質賃金（現金給与総額）", months=months, stats_code="00450071")
        return parse_real_wage_xls(_get(url, timeout=120)), {"file": url.replace(app_id, "***"), "catalog_month": ym}
    raise KeyError(kind)


def fetch_estat_series(kind: str, today=None) -> Points:
    return fetch_estat_file(kind, today)[0]


# The appId stays on the machine that has it. What that machine reads is written to a small cache file
# in the repo, so the scheduled run (which has no key) keeps the cards and refreshes them from the
# cache; only a run with the key moves the two series forward.
ESTAT_CACHE = Path(__file__).resolve().parent.parent / "config" / "jp_estat_series_v1.json"
ESTAT_CACHE_FROM = "2010-01-01"


def load_estat_cache(path: Path = ESTAT_CACHE) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {"schema_version": "jp-estat-series-v1", "series": {}}


def save_estat_cache(cache: dict[str, Any], updates: dict[str, tuple[Points, dict[str, str]]], now_iso: str, path: Path = ESTAT_CACHE) -> bool:
    """Merge fresh reads into the cache. The file is rewritten (and `retrieved_at` moved) only when a
    series actually changed. Returns True when it was written."""
    series = dict(cache.get("series") or {})
    for kind, (points, info) in updates.items():
        series[kind] = {**info, "points": [[d, round(v, 4)] for d, v in points if d >= ESTAT_CACHE_FROM]}
    if series == (cache.get("series") or {}):
        return False
    path.write_text(json.dumps({"schema_version": "jp-estat-series-v1", "retrieved_at": now_iso, "series": series},
                               ensure_ascii=False, indent=1, allow_nan=False) + "\n", encoding="utf-8")
    return True


def estat_reader(cache: dict[str, Any], updates: dict[str, tuple[Points, dict[str, str]]]) -> Callable[[str], Points]:
    """Live read when the key is set (and remembered in `updates`), the cached series when it is not."""
    def read(kind: str) -> Points:
        try:
            points, info = fetch_estat_file(kind)
        except MissingKey:
            cached = (cache.get("series") or {}).get(kind)
            if not cached:
                raise
            return [(d, float(v)) for d, v in cached["points"]]
        updates[kind] = (points, info)
        return points
    return read


# --------------------------------------------------------------------------
# Transforms
# --------------------------------------------------------------------------

def ratio_to_quarterly(monthly: Points, quarterly_base: Points, *, pct: bool = True) -> Points:
    """monthly / (the quarter's figure) for each month, from the first quarter that has one on. A
    month after the latest published quarter uses that latest quarter -- the most recent GDP, as the
    ratio is always quoted."""
    if not quarterly_base:
        raise ValueError("no quarterly base")
    by_q = dict(quarterly_base)
    last_q = quarterly_base[-1][0]
    first_q = quarterly_base[0][0]
    out: Points = []
    for d, v in monthly:
        qm = f"{d[:4]}-{((int(d[5:7]) - 1) // 3) * 3 + 1:02d}-01"
        if qm < first_q:
            continue
        base = by_q.get(qm, by_q[last_q] if qm > last_q else None)
        if base:
            out.append((d, v / base * (100 if pct else 1)))
    return out


def spread_bp(a: Points, b: Points) -> Points:
    bd = dict(b)
    return [(d, (v - bd[d]) * 100) for d, v in a if d in bd]


def pin_last_date(patch: dict[str, Any], last_obs: str) -> None:
    """History is drawn on a month-end grid; when the newest month is not over, its point is the
    observation's own date rather than a month-end that has not happened."""
    for w in patch["history"].values():
        if w["dates"] and last_obs[:7] == w["dates"][-1][:7] and last_obs < w["dates"][-1]:
            w["dates"][-1] = last_obs


# --------------------------------------------------------------------------
# What each indicator is made of
# --------------------------------------------------------------------------

def _urls(*u: str) -> tuple[str, ...]:
    return tuple(u)


BOJ_SRC = "boj:stat-search"
_BOJ_PAGE = "https://www.stat-search.boj.or.jp/"
_MOF_PAGE = "https://www.mof.go.jp/english/policy/jgbs/reference/interest_rate/"

SPECS: dict[str, Spec] = {s.id: s for s in [
    Spec("boj_total_assets", "monthly", "tn_jpy", "tn1", "BOJ 대차대조표 총자산(월말 잔액)입니다. 일본은행 시계열 통계 BS01(Bank of Japan Accounts / Total).",
         BOJ_SRC, _urls(_BOJ_PAGE), "monthly"),
    Spec("boj_assets_yoy", "monthly", "%", "pct1", "BOJ 총자산의 전년 같은 달 대비 증가율입니다. 위 총자산 시계열에서 계산했습니다.",
         BOJ_SRC, _urls(_BOJ_PAGE), "monthly"),
    Spec("boj_assets_gdp", "monthly", "%", "pct0", "BOJ 총자산 ÷ 명목 GDP(내각부, 연율). 분기 GDP가 아직 안 나온 달은 가장 최근 분기 GDP를 씁니다.",
         "boj+fred:JPNNGDP", _urls(_BOJ_PAGE, "https://fred.stlouisfed.org/series/JPNNGDP"), "monthly"),
    Spec("boj_etf", "monthly", "tn_jpy", "tn1", "BOJ가 신탁으로 보유한 지수연동 ETF(장부가, 시가 아님). BS01 'Pecuniary Trusts (Index-Linked ETFs)'.",
         BOJ_SRC, _urls(_BOJ_PAGE), "monthly"),
    Spec("boj_jreit", "monthly", "bn_jpy", "bn0u", "BOJ가 신탁으로 보유한 J-REIT(장부가). BS01 'Pecuniary Trusts (J-REITs)'.",
         BOJ_SRC, _urls(_BOJ_PAGE), "monthly"),
    Spec("call_rate", "monthly", "%", "pct2", "무담보 콜금리(익일물)의 월평균입니다. 일본은행 FM02.",
         BOJ_SRC, _urls(_BOJ_PAGE), "monthly"),
    Spec("m2_vs_2019", "monthly", "%", "pct1", "(현재 M2 평잔 − 2019-12) / 2019-12 × 100. 일본은행 MD02 M2 평잔(원계열) 기준입니다.",
         BOJ_SRC, _urls(_BOJ_PAGE), "monthly"),
    Spec("cgpi", "monthly", "%", "pct1", "기업물가지수(PPI, 옛 CGPI 2020년 기준) 총합의 전년 대비 상승률입니다. 일본은행이 CGPI를 PPI로 이름을 바꿨습니다.",
         BOJ_SRC, _urls(_BOJ_PAGE), "monthly", label_ko="기업물가 PPI(구 CGPI) YoY"),
    Spec("current_account", "monthly", "tn_jpy", "tn1s", "월별 경상수지 순잔액(조엔, 원계열)입니다. 일본은행 국제수지 BP01. 다른 나라 카드는 달러 기준이라 단위가 다릅니다.",
         BOJ_SRC, _urls(_BOJ_PAGE), "monthly", label_ko="경상수지(월, 조엔)"),
    Spec("gdp_yoy", "quarterly", "%", "pct1", "실질 GDP(내각부, 계절조정)의 전년 같은 분기 대비 증가율입니다. FRED JPNRGDPEXP에서 계산했습니다.",
         "fred:JPNRGDPEXP", _urls("https://fred.stlouisfed.org/series/JPNRGDPEXP"), "quarterly"),
    Spec("gdp_qoq", "quarterly", "%", "pct1", "실질 GDP의 직전 분기 대비 증가율입니다(연율 환산 아님). FRED JPNRGDPEXP에서 계산했습니다.",
         "fred:JPNRGDPEXP", _urls("https://fred.stlouisfed.org/series/JPNRGDPEXP"), "quarterly"),
    Spec("fx_reserves", "monthly", "bn_usd", "bn0u", "외환보유액(금 제외, 달러)입니다. FRED TRESEGJPM052N(IMF 기반).",
         "fred:TRESEGJPM052N", _urls("https://fred.stlouisfed.org/series/TRESEGJPM052N"), "monthly", label_ko="외환보유액(금 제외)"),
    Spec("bond_2y", "monthly", "%", "pct2", "일본 국채 2년물 수익률(재무성 공표, 월말 마지막 거래일; 이번 달은 최신일).",
         "mof:jgbcme", _urls(_MOF_PAGE), "daily"),
    Spec("bond_10y", "monthly", "%", "pct2", "일본 국채 10년물 수익률(재무성 공표, 월말 마지막 거래일; 이번 달은 최신일).",
         "mof:jgbcme", _urls(_MOF_PAGE), "daily"),
    Spec("bond_30y", "monthly", "%", "pct2", "일본 국채 30년물 수익률(재무성 공표, 월말 마지막 거래일; 이번 달은 최신일).",
         "mof:jgbcme", _urls(_MOF_PAGE), "daily"),
    Spec("spread_10y2y", "monthly", "bp", "bp0", "10년물 − 2년물 수익률(bp). 재무성 공표 수익률에서 같은 날짜끼리 뺐습니다.",
         "mof:jgbcme", _urls(_MOF_PAGE), "daily"),
    Spec("spread_30y10y", "monthly", "bp", "bp0", "30년물 − 10년물 수익률(bp). 재무성 공표 수익률에서 같은 날짜끼리 뺐습니다.",
         "mof:jgbcme", _urls(_MOF_PAGE), "daily"),
    Spec("core_cpi_jp", "monthly", "%", "pct1", "전국 소비자물가 '생선식품 제외 종합'(근원)의 전년 같은 달 대비 상승률입니다. 총무성 통계국 공표 지수에서 계산했습니다. 2025년 12월까지는 2020년 기준, 2026년 1월부터는 2025년 기준 지수를 씁니다(통계국이 기준 개편 후 그렇게 비교합니다).",
         "stat.go.jp:cpi", _urls("https://www.stat.go.jp/data/cpi/index.html"), "monthly"),
    Spec("core_core_cpi", "monthly", "%", "pct1", "전국 소비자물가 '생선식품 및 에너지 제외 종합'(근원-근원)의 전년 같은 달 대비 상승률입니다. 기준 이음은 근원 CPI와 같습니다.",
         "stat.go.jp:cpi", _urls("https://www.stat.go.jp/data/cpi/index.html"), "monthly"),
    Spec("tokyo_cpi", "monthly", "%", "pct1", "도쿄 23구 '생선식품 제외 종합'의 전년 같은 달 대비 상승률입니다(월간 확정치; 월말에 먼저 나오는 중순 속보가 아닙니다). 통계국이 2025년 기준 파일만 공개해서 2026년 1월부터의 값만 있습니다.",
         "stat.go.jp:cpi", _urls("https://www.stat.go.jp/data/cpi/index.html"), "monthly", label_ko="도쿄 근원 CPI(생선식품 제외) YoY"),
    Spec("job_applicant_ratio", "monthly", "ratio", "ratio2", "유효구인배율(계절조정, 파트 포함 일반)입니다. 후생노동성 일반직업소개상황 장기시계열표 3. " + ESTAT_CREDIT,
         "mhlw+e-stat", _urls("https://www.e-stat.go.jp/"), "monthly"),
    Spec("real_wage_yoy", "monthly", "%", "pct1", "실질임금 전년 동월 대비(현금급여총액, 조사산업계, 5인 이상, 소비자물가 '자가귀속임대료 제외 종합'으로 디플레이트)입니다. 후생노동성 매월근로통계 장기시계열표 25-1. 속보치는 나중에 개정됩니다. " + ESTAT_CREDIT,
         "mhlw+e-stat", _urls("https://www.e-stat.go.jp/"), "monthly", label_ko="실질임금 YoY(현금급여총액)"),
    Spec("yen_imm_net", "weekly", "k_contracts", "k0s", "CME 엔 선물의 비상업(투기) 순포지션 = 롱 − 숏(천 계약). CFTC 주간 보고서(화요일 기준, 금요일 공표).",
         "cftc:097741", _urls("https://www.cftc.gov/MarketReports/CommitmentsofTraders/index.htm"), "weekly"),
]}

FRED_IDS = ("JPNRGDPEXP", "JPNNGDP", "TRESEGJPM052N")
CPI_ITEMS = ("All items, less fresh food", "All items, less fresh food and energy")
CPI_REBASE = "2026-01-01"          # the 2025-base index gives the year-on-year from this month


class Sources:
    """Lazy, cached fetchers so one run reads each upstream series once."""

    def __init__(self, *, boj=fetch_boj, fred=ups.fetch_fred, mof=fetch_mof_jgb, cftc=fetch_cftc_yen, stat_cpi=fetch_stat_cpi,
                 estat=fetch_estat_series):
        self._boj, self._fred, self._mof, self._cftc, self._stat_cpi, self._estat = boj, fred, mof, cftc, stat_cpi, estat
        self._cache: dict[Any, Any] = {}

    def _memo(self, key, fn: Callable[[], Any]):
        if key not in self._cache:
            self._cache[key] = fn()
        return self._cache[key]

    def boj(self, db: str, code: str) -> Points:
        return self._memo(("boj", db, code), lambda: self._boj(db, code))

    def fred(self, sid: str) -> Points:
        return self._memo(("fred", sid), lambda: self._fred(sid))

    def jgb(self) -> dict[str, Points]:
        return self._memo("jgb", self._mof)

    def yen(self) -> Points:
        return self._memo("cftc", self._cftc)

    def estat(self, kind: str) -> Points:
        return self._memo(("estat", kind), lambda: self._estat(kind))

    def cpi(self, kind: str, base: int) -> dict[str, Points]:
        return self._memo(("cpi", kind, base), lambda: self._stat_cpi(kind, base, CPI_ITEMS))


def series_for(spec_id: str, s: Sources) -> Points:
    if spec_id == "boj_total_assets":
        return ups.scale(s.boj("BS01", "MABJMTA"), 1e-4)                       # 100 million yen -> trillion yen
    if spec_id == "boj_assets_yoy":
        return ups.pct_change(s.boj("BS01", "MABJMTA"), 12)
    if spec_id == "boj_assets_gdp":
        assets_tn = ups.scale(s.boj("BS01", "MABJMTA"), 1e-4)
        gdp_tn = ups.scale(s.fred("JPNNGDP"), 1e-3)                            # billion yen -> trillion yen
        return ratio_to_quarterly(assets_tn, gdp_tn)
    if spec_id == "boj_etf":
        return ups.scale(s.boj("BS01", "MABJMA003"), 1e-4)
    if spec_id == "boj_jreit":
        return ups.scale(s.boj("BS01", "MABJMA004"), 0.1)                      # 100 million yen -> billion yen
    if spec_id == "call_rate":
        return s.boj("FM02", "STRACLUCON")
    if spec_id == "m2_vs_2019":
        return ups.vs_base(s.boj("MD02", "MAM1NAM2M2MO"), "2019-12-01")
    if spec_id == "cgpi":
        return s.boj("PR01", "PRCG20_2200000000%")
    if spec_id == "current_account":
        return ups.scale(s.boj("BP01", "BPBP6JYNCB"), 1e-4)
    if spec_id == "gdp_yoy":
        return ups.pct_change(s.fred("JPNRGDPEXP"), 4)
    if spec_id == "gdp_qoq":
        return ups.pct_change(s.fred("JPNRGDPEXP"), 1)
    if spec_id == "fx_reserves":
        return ups.scale(s.fred("TRESEGJPM052N"), 1e-3)                        # million USD -> billion USD
    if spec_id in ("bond_2y", "bond_10y", "bond_30y"):
        return month_last(s.jgb()[spec_id.split("_")[1].upper()])
    if spec_id == "spread_10y2y":
        return spread_bp(month_last(s.jgb()["10Y"]), month_last(s.jgb()["2Y"]))
    if spec_id == "spread_30y10y":
        return spread_bp(month_last(s.jgb()["30Y"]), month_last(s.jgb()["10Y"]))
    if spec_id == "yen_imm_net":
        return ups.scale(s.yen(), 1e-3)
    if spec_id == "job_applicant_ratio":
        return s.estat("job_ratio")
    if spec_id == "real_wage_yoy":
        return s.estat("real_wage")
    if spec_id in ("core_cpi_jp", "core_core_cpi"):
        item = CPI_ITEMS[0] if spec_id == "core_cpi_jp" else CPI_ITEMS[1]
        return chained_yoy(s.cpi("zmi", 2020)[item], s.cpi("zmi", 2025)[item], switch=CPI_REBASE)
    if spec_id == "tokyo_cpi":
        tokyo = s.cpi("tmi", 2025)[CPI_ITEMS[0]]
        return chained_yoy([], tokyo, switch=CPI_REBASE)
    raise KeyError(spec_id)


DAILY_SOURCED = {"bond_2y", "bond_10y", "bond_30y", "spread_10y2y", "spread_30y10y"}


def build_patch(spec_id: str, points: Points, *, retrieved_at: str) -> dict[str, Any]:
    spec = SPECS[spec_id]
    patch = ups.build_patch(spec, points, retrieved_at=retrieved_at, asof=points[-1][0] if spec_id in DAILY_SOURCED else None)
    if spec_id in DAILY_SOURCED:
        pin_last_date(patch, points[-1][0])
    return patch


# --------------------------------------------------------------------------
# Applying to the pack
# --------------------------------------------------------------------------

def apply_gdp_composite(by_id: dict[str, dict[str, Any]], retrieved_at: str, source: str = "fred:JPNRGDPEXP") -> bool:
    """The 'gdp' chip is YoY | QoQ; rebuild it from the two real series (and keep its contribution
    breakdown, which is a separate view)."""
    gdp, yoy, qoq = by_id.get("gdp"), by_id.get("gdp_yoy"), by_id.get("gdp_qoq")
    if not (gdp and yoy and qoq and yoy.get("quality") == "live" and qoq.get("quality") == "live"):
        return False

    def mode(label: str, src: dict[str, Any]) -> dict[str, Any]:
        return {"label_ko": label, "id": src["id"], "value": src["value"], "display": src["display"], "unit": "%", "history": src["history"]}

    patch = {
        "value": yoy["value"], "display": f"{yoy['display']} | {qoq['display']}", "display_chip": f"{yoy['display']} | {qoq['display']}",
        "asof": yoy["asof"], "observed_at": yoy["observed_at"], "reference_period": yoy["reference_period"],
        "source": source, "source_urls": yoy["source_urls"], "quality": "live", "data_status": "live",
        "change_1m_pct": None, "change_1y_pct": None,
        "modes": {"yoy": mode("YoY", yoy), "qoq": {**mode("QoQ", qoq), "note_ko": "분기 대비 % (연율 환산 아님)"}},
        "retrieved_at": retrieved_at,
    }
    if any(gdp.get(k) != v for k, v in ups._data_of(patch).items()):
        gdp.update(patch)
        return True
    return False


def apply_etf_holdings(by_id: dict[str, dict[str, Any]], retrieved_at: str) -> bool:
    """'boj_etf_holdings' was the ETF amount plus the BOJ's share of the ETF market. The share has no free
    source, so the card now carries the real amount only -- not a real amount next to an invented share."""
    card, etf = by_id.get("boj_etf_holdings"), by_id.get("boj_etf")
    if not (card and etf and etf.get("quality") == "live"):
        return False
    patch = {
        "value": etf["value"], "display": etf["display"], "display_chip": etf["display"], "asof": etf["asof"], "observed_at": etf["observed_at"],
        "source": etf["source"], "source_urls": etf["source_urls"], "quality": "live", "data_status": "live", "derived": False,
        "components": [{"id": "boj_etf", "label_ko": "BOJ 보유 ETF(장부가)", "value": etf["value"], "display": etf["display"], "unit": "tn_jpy"}],
        "history": etf["history"], "note_ko": "BOJ 보유 ETF 장부가 잔액. 'ETF시장 내 비중'은 무료 출처가 없어 싣지 않습니다.",
        "change_1m_pct": None, "change_1y_pct": None, "retrieved_at": retrieved_at,
    }
    if any(card.get(k) != v for k, v in ups._data_of(patch).items()):
        card.update(patch)
        return True
    return False


def apply_all(jpn: dict[str, Any], patches: dict[str, dict[str, Any]], *, retrieved_at: str) -> dict[str, Any]:
    by_id = {i["id"]: i for i in jpn["indicators"]}
    changed = [k for k, p in patches.items() if k in by_id and ups.apply_patch(by_id[k], p)]
    if apply_gdp_composite(by_id, retrieved_at):
        changed.append("gdp")
    if apply_etf_holdings(by_id, retrieved_at):
        changed.append("boj_etf_holdings")
    ups.sync_chips(jpn, by_id, set(changed))
    before = dict(jpn.get("data_status_summary") or {})
    ups.refresh_status_summary(jpn)
    return {"changed": changed, "summary_changed": before != jpn["data_status_summary"]}
