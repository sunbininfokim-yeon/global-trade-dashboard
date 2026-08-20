"""Official-source builder for the Japan growth-monitor additions.

The dashboard build itself is deliberately offline and deterministic.  This
module is the explicit refresh boundary: it downloads the three primary-source
files, constructs the two ratios, and writes a small versioned snapshot that
the normal macro builder can consume without a network dependency.
"""

from __future__ import annotations

import csv
import io
import json
import re
from collections.abc import Iterable
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import load_workbook


BOJ_API_URL = "https://www.stat-search.boj.or.jp/api/v1/getDataCode"
BOJ_OUTPUT_GAP_URL = "https://www.boj.or.jp/en/research/research_data/gap/gap.xlsx"

# ESRI publishes a new directory every quarterly-release vintage.  Keep the
# release page explicit so a review can reproduce the exact Cabinet Office
# vintage rather than silently switching revisions during a build.
ESRI_RELEASE_PAGE_URL = (
    "https://www.esri.cao.go.jp/en/sna/data/sokuhou/files/2026/qe261_2/gdemenuea.html"
)

PRIVATE_NONFINANCIAL_FFA = "FOF_FFAF411L700"
GENERAL_GOVERNMENT_FFA = "FOF_FFAF420L700"


def fetch_bytes(url: str, *, timeout: int = 60) -> bytes:
    """Fetch a public source with a descriptive user agent."""
    request = Request(url, headers={"User-Agent": "global-trade-dashboard/1.0"})
    with urlopen(request, timeout=timeout) as response:  # nosec B310 - fixed public source URLs
        return response.read()


def _number(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip().replace(",", "")
    if not text or text in {"-", "…", "..."}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _quarter_end(period: int | str) -> str:
    """Turn BOJ YYYYQQ periods into ISO quarter-end dates."""
    text = str(period)
    if not re.fullmatch(r"\d{6}", text):
        raise ValueError(f"not a YYYYQQ period: {period!r}")
    year, quarter = int(text[:4]), int(text[-2:])
    if quarter not in (1, 2, 3, 4):
        raise ValueError(f"invalid quarter: {period!r}")
    return f"{year:04d}-{(3, 6, 9, 12)[quarter - 1]:02d}-{(31, 30, 30, 31)[quarter - 1]:02d}"


def _quarter_key(day: str) -> tuple[int, int]:
    parsed = date.fromisoformat(day)
    return parsed.year, (parsed.month - 1) // 3 + 1


def _previous_quarter(key: tuple[int, int]) -> tuple[int, int]:
    year, quarter = key
    return (year - 1, 4) if quarter == 1 else (year, quarter - 1)


def _is_four_quarter_window(dates: Iterable[str]) -> bool:
    keys = [_quarter_key(day) for day in dates]
    return len(keys) == 4 and all(keys[i - 1] == _previous_quarter(keys[i]) for i in range(1, 4))


def _decode_esri_csv(payload: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "cp932", "shift_jis"):
        try:
            return list(csv.reader(io.StringIO(payload.decode(encoding))))
        except UnicodeDecodeError:
            continue
    raise ValueError("ESRI CSV could not be decoded as UTF-8 or CP932")


def parse_esri_nominal_calendar_year(payload: bytes) -> list[dict[str, Any]]:
    """Return annual nominal private non-residential investment / GDP ratios."""
    rows = _decode_esri_csv(payload)
    header_index = next(
        (
            index
            for index, row in enumerate(rows)
            if "GDP(Expenditure Approach)" in row and "Private Non-Resi.Investment" in row
        ),
        None,
    )
    if header_index is None:
        raise ValueError("ESRI annual nominal GDP columns were not found")
    header = rows[header_index]
    gdp_col = header.index("GDP(Expenditure Approach)")
    capex_col = header.index("Private Non-Resi.Investment")
    observations: list[dict[str, Any]] = []
    for row in rows[header_index + 1 :]:
        if not row or not re.match(r"^\d{4}/", (row[0] or "").strip()):
            continue
        if max(gdp_col, capex_col) >= len(row):
            continue
        gdp, capex = _number(row[gdp_col]), _number(row[capex_col])
        if not gdp or capex is None:
            continue
        year = int(row[0][:4])
        observations.append(
            {
                "date": f"{year:04d}-12-31",
                "value": round(capex / gdp * 100.0, 4),
                "gdp_bn_jpy": round(gdp, 1),
                "private_non_resi_investment_bn_jpy": round(capex, 1),
            }
        )
    if not observations:
        raise ValueError("ESRI annual nominal GDP contained no observations")
    return observations


def parse_esri_nominal_quarterly_gdp(payload: bytes) -> dict[str, float]:
    """Read the original (not annualised) nominal GDP quarter values in bn JPY."""
    rows = _decode_esri_csv(payload)
    header_index = next(
        (index for index, row in enumerate(rows) if "GDP(Expenditure Approach)" in row),
        None,
    )
    if header_index is None:
        raise ValueError("ESRI quarterly nominal GDP column was not found")
    gdp_col = rows[header_index].index("GDP(Expenditure Approach)")
    current_year: int | None = None
    out: dict[str, float] = {}
    for row in rows[header_index + 1 :]:
        if not row or gdp_col >= len(row):
            continue
        period = (row[0] or "").strip()
        year_match = re.match(r"^(\d{4})/\s*(\d{1,2})-\s*(\d{1,2})\.", period)
        if year_match:
            current_year = int(year_match.group(1))
            start_month, end_month = int(year_match.group(2)), int(year_match.group(3))
        else:
            range_match = re.match(r"^(\d{1,2})-\s*(\d{1,2})\.", period)
            if not range_match or current_year is None:
                continue
            start_month, end_month = int(range_match.group(1)), int(range_match.group(2))
        quarter = {3: 1, 6: 2, 9: 3, 12: 4}.get(end_month)
        gdp = _number(row[gdp_col])
        if quarter is None or gdp is None:
            continue
        out[_quarter_end(f"{current_year:04d}{quarter:02d}")] = gdp
    if not out:
        raise ValueError("ESRI quarterly nominal GDP contained no observations")
    return out


def parse_boj_output_gap(payload: bytes) -> list[dict[str, Any]]:
    """Read BOJ's quarterly output gap (% of potential GDP) from gap.xlsx."""
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        worksheet = workbook["data1"]
    except KeyError as exc:
        raise ValueError("BOJ output-gap workbook has no data1 sheet") from exc
    observations: list[dict[str, Any]] = []
    for row in worksheet.iter_rows(min_row=1, values_only=True):
        quarter = row[0] if row else None
        if not isinstance(quarter, str):
            continue
        match = re.fullmatch(r"(\d{4})\.(\d)Q", quarter.strip())
        if not match:
            continue
        value = _number(row[1] if len(row) > 1 else None)
        if value is None:
            continue
        observations.append(
            {"date": _quarter_end(f"{match.group(1)}0{match.group(2)}"), "value": round(value, 4)}
        )
    if not observations:
        raise ValueError("BOJ output-gap workbook contained no observations")
    return observations


def parse_boj_api_series(payload: bytes) -> dict[str, dict[str, float | None]]:
    """Map BOJ API result rows to ``series code -> ISO date -> numeric value``."""
    doc = json.loads(payload.decode("utf-8"))
    if doc.get("STATUS") != 200:
        raise ValueError(f"BOJ API failed: {doc.get('MESSAGE') or doc.get('STATUS')}")
    out: dict[str, dict[str, float | None]] = {}
    for row in doc.get("RESULTSET") or []:
        values = row.get("VALUES") or {}
        periods = values.get("SURVEY_DATES") or []
        raw_values = values.get("VALUES") or []
        if len(periods) != len(raw_values):
            raise ValueError(f"BOJ API length mismatch for {row.get('SERIES_CODE')}")
        out[str(row["SERIES_CODE"])] = {
            _quarter_end(period): _number(value) for period, value in zip(periods, raw_values)
        }
    return out


def build_net_funding_demand(
    private_nonfinancial: dict[str, float | None],
    general_government: dict[str, float | None],
    quarterly_gdp_bn_jpy: dict[str, float],
) -> list[dict[str, Any]]:
    """Construct 4-quarter net funding demand as a percentage of nominal GDP.

    BOJ's financial surplus/deficit is a quarterly flow in 100 million JPY.  A
    four-quarter sum of each sector is therefore divided by the matching
    four-quarter sum of ESRI nominal GDP (bn JPY); 100 million JPY is 0.1 bn.
    Positive values mean the combined corporate-plus-government sector is a
    net financial surplus; negative values mean it is a net funding demander.
    """
    common = sorted(
        day
        for day in quarterly_gdp_bn_jpy
        if private_nonfinancial.get(day) is not None and general_government.get(day) is not None
    )
    observations: list[dict[str, Any]] = []
    for index in range(3, len(common)):
        window = common[index - 3 : index + 1]
        if not _is_four_quarter_window(window):
            continue
        corporate = sum(float(private_nonfinancial[day]) for day in window)
        fiscal = sum(float(general_government[day]) for day in window)
        gdp = sum(float(quarterly_gdp_bn_jpy[day]) for day in window)
        if not gdp:
            continue
        corporate_pct = corporate * 0.1 / gdp * 100.0
        fiscal_pct = fiscal * 0.1 / gdp * 100.0
        observations.append(
            {
                "date": window[-1],
                "value": round(corporate_pct + fiscal_pct, 4),
                "private_nonfinancial_pct_gdp": round(corporate_pct, 4),
                "general_government_pct_gdp": round(fiscal_pct, 4),
                "nominal_gdp_bn_jpy_4q": round(gdp, 1),
            }
        )
    if not observations:
        raise ValueError("no four-quarter BOJ FFA / ESRI GDP overlap")
    return observations


def _esri_table_urls(release_page_url: str, release_page: bytes) -> tuple[str, str]:
    """Find the calendar-year and quarterly nominal GDP CSVs on an ESRI page."""
    from urllib.parse import urljoin

    html = release_page.decode("utf-8", errors="replace")
    annual = re.search(r'href="([^"]*/gaku-mcy\d+\.csv)"', html)
    quarterly = re.search(r'href="([^"]*/gaku-mg\d+\.csv)"', html)
    if not annual or not quarterly:
        raise ValueError("ESRI release page did not contain nominal GDP CSV links")
    return urljoin(release_page_url, annual.group(1)), urljoin(release_page_url, quarterly.group(1))


def _boj_api_url(codes: list[str], *, start_date: str, end_date: str | None = None) -> str:
    params: list[tuple[str, str]] = [
        ("format", "JSON"),
        ("lang", "EN"),
        ("db", "FF"),
        ("startDate", start_date),
        ("code", ",".join(codes)),
    ]
    if end_date:
        params.insert(5, ("endDate", end_date))
    # BOJ's API accepts a comma-separated code list but rejects an encoded
    # comma (%2C) with HTTP 400.
    return f"{BOJ_API_URL}?{urlencode(params, safe=',')}"


def build_snapshot(
    *,
    release_page_url: str = ESRI_RELEASE_PAGE_URL,
    retrieved_at: str | None = None,
    fetcher: Any = fetch_bytes,
) -> dict[str, Any]:
    """Fetch official inputs and return the committed Japan-growth snapshot."""
    release_page = fetcher(release_page_url)
    annual_url, quarterly_url = _esri_table_urls(release_page_url, release_page)
    calendar_year = parse_esri_nominal_calendar_year(fetcher(annual_url))
    quarterly_gdp = parse_esri_nominal_quarterly_gdp(fetcher(quarterly_url))
    ffa_url = _boj_api_url(
        [PRIVATE_NONFINANCIAL_FFA, GENERAL_GOVERNMENT_FFA],
        start_date="199801",
    )
    ffa = parse_boj_api_series(fetcher(ffa_url))
    net_funding = build_net_funding_demand(
        ffa[PRIVATE_NONFINANCIAL_FFA],
        ffa[GENERAL_GOVERNMENT_FFA],
        quarterly_gdp,
    )
    output_gap = parse_boj_output_gap(fetcher(BOJ_OUTPUT_GAP_URL))
    stamp = retrieved_at or datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")

    latest_net = net_funding[-1]
    return {
        "schema_version": "jp-growth-monitor-v1",
        "retrieved_at": stamp,
        "source": {
            "cabinet_office_release_page": release_page_url,
            "cabinet_office_nominal_calendar_year_csv": annual_url,
            "cabinet_office_nominal_quarterly_csv": quarterly_url,
            "boj_output_gap_xlsx": BOJ_OUTPUT_GAP_URL,
            "boj_ffa_api": ffa_url,
        },
        "series": {
            "capex_gdp_ratio": {
                "frequency": "annual",
                "asof": calendar_year[-1]["date"],
                "unit": "pct_gdp",
                "formula_ko": "명목 민간기업 설비투자 ÷ 명목 GDP × 100",
                "observations": calendar_year,
            },
            "net_funding_demand": {
                "frequency": "quarterly_rolling_4q",
                "asof": latest_net["date"],
                "unit": "pct_gdp",
                "formula_ko": (
                    "(BOJ 민간비금융법인 금융잉여/부족 + 일반정부 금융잉여/부족)의 "
                    "4분기 합계 ÷ 같은 4분기 명목 GDP 합계 × 100"
                ),
                "sign_convention_ko": "음수=기업·정부 합산 순자금수요(투자초과), 양수=순자금잉여.",
                "components_latest": {
                    "private_nonfinancial_pct_gdp": latest_net["private_nonfinancial_pct_gdp"],
                    "general_government_pct_gdp": latest_net["general_government_pct_gdp"],
                    "nominal_gdp_bn_jpy_4q": latest_net["nominal_gdp_bn_jpy_4q"],
                },
                "observations": net_funding,
            },
            "gdp_gap": {
                "frequency": "quarterly",
                "asof": output_gap[-1]["date"],
                "unit": "pct_potential_gdp",
                "formula_ko": "BOJ 추정 GDP 갭 (% of potential GDP)",
                "observations": output_gap,
            },
        },
    }


def write_snapshot(path: Path, snapshot: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
