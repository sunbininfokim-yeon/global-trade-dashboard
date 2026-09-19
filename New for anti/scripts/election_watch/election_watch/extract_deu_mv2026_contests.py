"""Build elections_contest_v1 for deu-2026-mv-landtag.

SMD + list heads: LAiV Statistisches Wahlheft 4/2026 XLSX (402 candidates).
Current seats: Landtag MV Plenum page, Stand 1. Juni 2026 (79 members).
"""
from __future__ import annotations

import json
import re
import urllib.request
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "raw" / "deu"
RAW_XLSX = RAW_DIR / "wahlheft4_2026.xlsx"
PUBLIC = ROOT.parents[1] / "public" / "data"
CONTESTS_DIR = PUBLIC / "elections_contests"

AS_OF = date.today().isoformat()
EVENT_ID = "deu-2026-mv-landtag"

XLSX = "https://www.laiv-mv.de/serviceassistent/download?id=1692101"
PDF = "https://www.laiv-mv.de/serviceassistent/download?id=1692097"
LAIV = "https://www.laiv-mv.de/Wahlen/Landtagswahlen/2026"
PLENUM = "https://www.landtag-mv.de/landtag/grundsaetzliches/plenum"
FRACTIONS = "https://www.landtag-mv.de/landtag/grundsaetzliches/parlamentarische-gremien/fraktionen"
LANDTAG_HOME = "https://www.landtag-mv.de/"
WIKI_DE = "https://de.wikipedia.org/wiki/Landtagswahl_in_Mecklenburg-Vorpommern_2026"
LANDTAG_WAHL = "https://www.landtag-mv.de/landtag/rund-um-wahlen/landtagswahl-2026"

UA = "global-trade-dashboard election_watch/deu-mv2026 (+research)"

WK_META: Dict[int, Tuple[str, str, str]] = {
    1: ("Greifswald", "그라이프스발트", "그라이프스발트"),
    2: ("Neubrandenburg I", "노이브란덴부르크 1구", "노이브란덴부르크"),
    3: ("Neubrandenburg II", "노이브란덴부르크 2구", "노이브란덴부르크"),
    4: ("Hansestadt Rostock I", "로스토크 1구", "로스토크"),
    5: ("Hansestadt Rostock II", "로스토크 2구", "로스토크"),
    6: ("Hansestadt Rostock III", "로스토크 3구", "로스토크"),
    7: ("Hansestadt Rostock IV", "로스토크 4구", "로스토크"),
    8: ("Schwerin I", "슈베린 1구", "슈베린"),
    9: ("Schwerin II", "슈베린 2구", "슈베린"),
    10: ("Wismar", "비스마르", "비스마르"),
    11: ("Landkreis Rostock I", "로스토크군 1구", "로스토크군"),
    12: ("Landkreis Rostock II", "로스토크군 2구", "로스토크군"),
    13: (
        "Mecklenburgische Seenplatte I - Vorpommern-Greifswald I",
        "메클렌부르크호수지구 1구·포어포메른-그라이프스발트 1구",
        "메클렌부르크호수지구",
    ),
    14: ("Mecklenburgische Seenplatte II", "메클렌부르크호수지구 2구", "메클렌부르크호수지구"),
    15: ("Landkreis Rostock III", "로스토크군 3구", "로스토크군"),
    16: ("Landkreis Rostock IV", "로스토크군 4구", "로스토크군"),
    17: ("Ludwigslust-Parchim I", "루트비히슬루스트-파르힘 1구", "루트비히슬루스트-파르힘"),
    18: ("Ludwigslust-Parchim II", "루트비히슬루스트-파르힘 2구", "루트비히슬루스트-파르힘"),
    19: ("Ludwigslust-Parchim III", "루트비히슬루스트-파르힘 3구", "루트비히슬루스트-파르힘"),
    20: ("Mecklenburgische Seenplatte III", "메클렌부르크호수지구 3구", "메클렌부르크호수지구"),
    21: ("Mecklenburgische Seenplatte IV", "메클렌부르크호수지구 4구", "메클렌부르크호수지구"),
    22: ("Mecklenburgische Seenplatte V", "메클렌부르크호수지구 5구", "메클렌부르크호수지구"),
    23: ("Vorpommern-Rügen I", "포어포메른-뤼겐 1구", "포어포메른-뤼겐"),
    24: ("Vorpommern-Rügen II - Stralsund III", "포어포메른-뤼겐 2구·슈트랄준트 3구", "포어포메른-뤼겐"),
    25: ("Vorpommern-Rügen III - Stralsund I", "포어포메른-뤼겐 3구·슈트랄준트 1구", "포어포메른-뤼겐"),
    26: ("Stralsund II", "슈트랄준트 2구", "슈트랄준트"),
    27: ("Nordwestmecklenburg I", "노르트베스트메클렌부르크 1구", "노르트베스트메클렌부르크"),
    28: ("Nordwestmecklenburg II", "노르트베스트메클렌부르크 2구", "노르트베스트메클렌부르크"),
    29: ("Vorpommern-Greifswald II", "포어포메른-그라이프스발트 2구", "포어포메른-그라이프스발트"),
    30: ("Vorpommern-Greifswald III", "포어포메른-그라이프스발트 3구", "포어포메른-그라이프스발트"),
    31: ("Ludwigslust-Parchim IV", "루트비히슬루스트-파르힘 4구", "루트비히슬루스트-파르힘"),
    32: ("Ludwigslust-Parchim V", "루트비히슬루스트-파르힘 5구", "루트비히슬루스트-파르힘"),
    33: ("Vorpommern-Rügen IV", "포어포메른-뤼겐 4구", "포어포메른-뤼겐"),
    34: ("Vorpommern-Rügen V", "포어포메른-뤼겐 5구", "포어포메른-뤼겐"),
    35: ("Vorpommern-Greifswald IV", "포어포메른-그라이프스발트 4구", "포어포메른-그라이프스발트"),
    36: ("Vorpommern-Greifswald V", "포어포메른-그라이프스발트 5구", "포어포메른-그라이프스발트"),
}

PARTY_FROM_XLSX = {
    "SPD": "SPD",
    "AfD": "AfD",
    "CDU": "CDU",
    "Die Linke": "LINKE",
    "GRÜNE": "GRÜNE",
    "FDP": "FDP",
    "Tierschutzpartei": "Tierschutz",
    "FREIE WÄHLER": "FW",
    "Die PARTEI": "PARTEI",
    "PIRATEN": "PIRATEN",
    "ÖDP": "ÖDP",
    "Bündnis C": "BuendnisC",
    "BSW": "BSW",
    "Handwerker Partei Deutschland": "Handwerker",
    "KPD": "KPD",
    "PdF": "PdF",
    "Team Freiheit": "TF",
    "Volt": "Volt",
    "WLD": "WLD",
    "LfK": "LfK",
    "Einzelbewerber": "IND",
    "Einzelbewerberin": "IND",
}

PARTY_META: Dict[str, Dict[str, str]] = {
    "SPD": {
        "name_de": "Sozialdemokratische Partei Deutschlands",
        "name_ko": "사민당",
        "name_en": "Social Democratic Party",
        "leader_ko": "마누엘라 슈베지히",
        "note_ko": "투표용지 1번. 주지사 후보·주명부 1번 슈베지히. 현 여당(좌파당과 연립).",
    },
    "AfD": {
        "name_de": "Alternative für Deutschland",
        "name_ko": "독일을 위한 대안",
        "name_en": "Alternative for Germany",
        "leader_ko": "엔리코 슐트",
        "note_ko": "투표용지 2번. 주명부 1번은 슐트. 주지사 후보는 연방하원 홀름(슈베린 1구 직후보지).",
    },
    "CDU": {
        "name_de": "Christlich Demokratische Union Deutschlands",
        "name_ko": "기민련",
        "name_en": "Christian Democratic Union",
        "leader_ko": "다니엘 페터스",
        "note_ko": "투표용지 3번. 주명부 1번 페터스.",
    },
    "LINKE": {
        "name_de": "Die Linke",
        "name_ko": "좌파당",
        "name_en": "The Left",
        "leader_ko": "시모네 올덴부르크",
        "note_ko": "투표용지 4번. 주명부 1번 올덴부르크. 현 연립 여당.",
    },
    "GRÜNE": {
        "name_de": "Bündnis 90/Die Grünen",
        "name_ko": "동맹90/녹색당",
        "name_en": "Alliance 90/The Greens",
        "leader_ko": "클라우디아 뮐러·올레 크뤼거",
        "note_ko": "투표용지 5번. 주명부 1번 뮐러, 2번 크뤼거.",
    },
    "FDP": {
        "name_de": "Freie Demokratische Partei",
        "name_ko": "자민당",
        "name_en": "Free Democratic Party",
        "leader_ko": "야코프 시르머",
        "note_ko": "투표용지 6번. 주명부 1번 시르머. 제8기에는 회파가 아니라 의회 그룹 3명.",
    },
    "Tierschutz": {
        "name_de": "PARTEI MENSCH KLIMA TIERSCHUTZ",
        "name_ko": "동물보호당",
        "name_en": "Human Environment Animal Protection",
        "leader_ko": "페트라 멜힌",
        "note_ko": "투표용지 7번. 주명부만(지역구 없음).",
    },
    "FW": {
        "name_de": "FREIE WÄHLER",
        "name_ko": "자유유권자",
        "name_en": "Free Voters",
        "leader_ko": "카를 케스너",
        "note_ko": "투표용지 8번. 주명부 1번 케스너.",
    },
    "PARTEI": {
        "name_de": "Die PARTEI",
        "name_ko": "디 파르타이",
        "name_en": "Die PARTEI",
        "leader_ko": "레아 알렉산드라 지베르트",
        "note_ko": "투표용지 9번. 주명부 1번 지베르트.",
    },
    "PIRATEN": {
        "name_de": "Piratenpartei Deutschland",
        "name_ko": "해적당",
        "name_en": "Pirate Party Germany",
        "leader_ko": "데니스 클뤼버",
        "note_ko": "투표용지 10번. 주명부만(지역구 없음).",
    },
    "ÖDP": {
        "name_de": "Ökologisch-Demokratische Partei",
        "name_ko": "생태민주당",
        "name_en": "Ecological Democratic Party",
        "leader_ko": "로날트 쉬네만",
        "note_ko": "투표용지 11번. 주명부만(지역구 없음).",
    },
    "BuendnisC": {
        "name_de": "Bündnis C - Christen für Deutschland",
        "name_ko": "동맹 C",
        "name_en": "Alliance C – Christians for Germany",
        "leader_ko": "크리스티안 하우저",
        "note_ko": "투표용지 12번. 주명부만(지역구 없음).",
    },
    "BSW": {
        "name_de": "Bündnis Sahra Wagenknecht – Vernunft und Gerechtigkeit",
        "name_ko": "사라 바겐크네히트 연합",
        "name_en": "Sahra Wagenknecht Alliance",
        "leader_ko": "페터 샤벨",
        "note_ko": "투표용지 13번. 주명부 1번 샤벨.",
    },
    "Handwerker": {
        "name_de": "Handwerker Partei Deutschland",
        "name_ko": "수공업자당",
        "name_en": "Craftsmen Party Germany",
        "leader_ko": "지크프리트 클라인",
        "note_ko": "투표용지 14번. 공식 약칭 없음. 주명부만(지역구 없음).",
    },
    "KPD": {
        "name_de": "Kommunistische Partei Deutschlands",
        "name_ko": "독일공산당",
        "name_en": "Communist Party of Germany",
        "leader_ko": "크리스티네 멜허",
        "note_ko": "투표용지 15번. 주명부만(지역구 없음).",
    },
    "PdF": {
        "name_de": "Partei des Fortschritts",
        "name_ko": "진보당",
        "name_en": "Party of Progress",
        "leader_ko": "수잔네 그라프",
        "note_ko": "투표용지 16번. 주명부만(지역구 없음).",
    },
    "TF": {
        "name_de": "Team Freiheit",
        "name_ko": "팀 자유",
        "name_en": "Team Freiheit",
        "leader_ko": "파울 브레셀",
        "note_ko": "투표용지 17번. 주명부 1번 브레셀. 위키 2026 총선 표는 제8기 의석 1을 여기로 잡지만, 주의회 공식 회파 표(2026-06-01)에는 무회파로만 나온다.",
    },
    "Volt": {
        "name_de": "Volt Deutschland",
        "name_ko": "볼트",
        "name_en": "Volt Germany",
        "leader_ko": "리자 리커",
        "note_ko": "투표용지 18번. 주명부 1번 리커.",
    },
    "WLD": {
        "name_de": "WIR LEBEN DEMOKRATIE",
        "name_ko": "우리는 민주주의를 산다",
        "name_en": "We Live Democracy",
        "leader_ko": "마르코 렉신",
        "note_ko": "투표용지 19번. 주명부 1번 렉신.",
    },
    "LfK": {
        "name_de": "Lobbyisten für Kinder",
        "name_ko": "아이들을 위한 로비스트",
        "name_en": "Lobbyists for Children",
        "note_ko": "주명부 없음. Wahlkreis 24(포어포메른-뤼겐 2구·슈트랄준트 3구) 지역구만.",
    },
    "IND": {
        "name_de": "Einzelbewerber",
        "name_ko": "무소속",
        "name_en": "Independent",
        "note_ko": "어느 정당·연립에도 합산하지 않는다. 제8기 무회파 4명도 여기(공식 회파 표 Stand 2026-06-01).",
    },
}

CURRENT_BY_PARTY = {
    "SPD": 34,
    "AfD": 13,
    "CDU": 12,
    "LINKE": 9,
    "GRÜNE": 4,
    "FDP": 3,
    "IND": 4,
}
CURRENT_TOTAL = 79
CURRENT_NOTE = (
    "제8기 메클렌부르크포어포메른 주의회 전체 79석(법정 71 + 초과의석·보정의석). "
    "landtag-mv.de Plenum 회파 Stand 2026-06-01: SPD 34, AfD 13(당선 14에서 1명 이탈), "
    "CDU 12, LINKE 9, GRÜNE 4, FDP는 회파가 아니라 의회 그룹 3명, 무회파 4. "
    "34+13+12+9+4+3+4=79. "
    "위키 2026 총선 표 'Mandate vor der Wahl 2026'는 AfD 14·Team Freiheit 1로 달라 "
    "공식 숫자를 쓰고 차이는 병기한다. "
    "소선거구·명부로 나눈 공시전 의석은 출처에 없어 두 열에 같은 전체 숫자를 둔다."
)

WK_HEADER_RE = re.compile(r"Wahlkreis:\s*(\d+)\s*-\s*(.+)$")
INCUMBENT_RE = re.compile(
    r"Abgeordnete(?:r|n)? des Landtages|Mitglied(?:es)? des Landtages|"
    r"\bMdL\b|Landtagsabgeordnet",
    re.I,
)


def compact(obj: Any) -> Any:
    if isinstance(obj, dict):
        return {k: compact(v) for k, v in obj.items() if v is not None and v != ""}
    if isinstance(obj, list):
        return [compact(x) for x in obj]
    return obj


def ensure_xlsx() -> Path:
    if RAW_XLSX.exists() and RAW_XLSX.stat().st_size > 10_000:
        return RAW_XLSX
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    req = urllib.request.Request(XLSX, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=60) as resp:
        data = resp.read()
    if len(data) < 10_000:
        raise SystemExit(f"xlsx too small: {len(data)} bytes from {XLSX}")
    RAW_XLSX.write_bytes(data)
    return RAW_XLSX


def party_abbr(raw: str) -> str:
    key = re.sub(r"\s+", " ", (raw or "").replace("\n", " ")).strip()
    if key in PARTY_FROM_XLSX:
        return PARTY_FROM_XLSX[key]
    raise SystemExit(f"unmapped party label: {raw!r}")


def format_name(family: Any, given: Any) -> str:
    fam = re.sub(r"\s+", " ", str(family or "")).strip()
    giv = re.sub(r"\s+", " ", str(given or "")).strip()
    if not fam:
        raise SystemExit(f"empty family name given={given!r}")
    return f"{fam}, {giv}" if giv else fam


def parse_smd(wb) -> List[Dict[str, Any]]:
    ws = wb["4. Kreiswahlvor. n. Bewerbern"]
    by_wk: Dict[int, List[Dict[str, Any]]] = defaultdict(list)
    wk_no: Optional[int] = None
    wk_name: Optional[str] = None
    for row in ws.iter_rows(values_only=True):
        label = row[1]
        if isinstance(label, str) and "Wahlkreis:" in label:
            m = WK_HEADER_RE.search(re.sub(r"\s+", " ", label).strip())
            if not m:
                raise SystemExit(f"bad WK header: {label!r}")
            wk_no = int(m.group(1))
            wk_name = m.group(2).strip()
            expected = WK_META[wk_no][0]
            if wk_name != expected:
                raise SystemExit(f"WK {wk_no} name {wk_name!r} != {expected!r}")
            continue
        if not isinstance(row[0], int) or not isinstance(label, str) or not row[3]:
            continue
        if wk_no is None:
            raise SystemExit("candidate row before WK header")
        party = party_abbr(label)
        cand: Dict[str, Any] = {
            "party_abbr": party,
            "name": format_name(row[3], row[4]),
            "status": "nominated",
            "source": XLSX,
        }
        if isinstance(row[2], int):
            cand["dual_listed"] = True
        job = str(row[5] or "")
        if INCUMBENT_RE.search(job):
            cand["incumbent"] = True
        by_wk[wk_no].append(cand)
    if set(by_wk) != set(range(1, 37)):
        raise SystemExit(f"expected WK 1-36, got {sorted(by_wk)}")
    out = []
    for n in range(1, 37):
        name_de, name_ko, group_ko = WK_META[n]
        cands = by_wk[n]
        if not cands:
            raise SystemExit(f"WK {n}: no candidates")
        out.append(
            {
                "id": f"{n:02d}",
                "name_ko": name_ko,
                "name_de": name_de,
                "group_ko": group_ko,
                "seats": 1,
                "status": "nominated",
                "candidates": cands,
            }
        )
    n_cands = sum(len(d["candidates"]) for d in out)
    if n_cands != 263:
        raise SystemExit(f"expected 263 SMD candidates, got {n_cands}")
    return out


def parse_pr_heads(wb) -> List[Dict[str, Any]]:
    ws = wb["2. Landeslisten nach Bewerbern"]
    heads: List[Dict[str, Any]] = []
    party: Optional[str] = None
    for row in ws.iter_rows(values_only=True):
        label = row[1]
        if isinstance(label, str) and re.match(r"\s*2\.\d+", label):
            text = re.sub(r"\s+", " ", label).strip()
            party = None
            for raw, abbr in PARTY_FROM_XLSX.items():
                if raw == "Einzelbewerber" or raw == "Einzelbewerberin" or raw == "LfK":
                    continue
                if raw in text or (raw == "Handwerker Partei Deutschland" and "Handwerker" in text):
                    party = abbr
            if party is None:
                raise SystemExit(f"unmapped list header: {text!r}")
            continue
        if row[0] != 1 or not isinstance(label, str) or party is None:
            continue
        heads.append(
            {
                "party_abbr": party,
                "name": format_name(label, row[2]),
                "status": "nominated",
                "source": XLSX,
                "note_ko": "주명부 1번.",
            }
        )
    if len(heads) != 19:
        raise SystemExit(f"expected 19 list heads, got {len(heads)} {[h['party_abbr'] for h in heads]}")
    seen = [h["party_abbr"] for h in heads]
    if len(set(seen)) != 19:
        raise SystemExit(f"duplicate list heads: {seen}")
    return heads


def pr_district(heads: List[Dict[str, Any]]) -> Dict[str, Any]:
    return {
        "id": "pr-mv",
        "name_ko": "주명부",
        "name_de": "Landeslisten",
        "group_ko": "비례",
        "seats": 35,
        "status": "nominated",
        "note_ko": (
            "명부 전원(313명)은 넣지 않고 19개 주명부 1번만 적는다. "
            "법정 명부 의석 35. 초과의석·보정의석으로 전체 71보다 커질 수 있다. "
            "LfK는 주명부가 없다."
        ),
        "candidates": heads,
    }


def parties_block(used: set) -> List[Dict[str, Any]]:
    out = []
    for abbr, meta in PARTY_META.items():
        if abbr not in used:
            continue
        row = {"abbr": abbr}
        for k in ("name_de", "name_ko", "name_en", "leader_ko", "note_ko"):
            if meta.get(k):
                row[k] = meta[k]
        out.append(row)
    missing = used - {p["abbr"] for p in out}
    if missing:
        raise SystemExit(f"PARTY_META missing {missing}")
    return out


def current_block() -> Dict[str, Any]:
    return {
        "by_party": dict(CURRENT_BY_PARTY),
        "total": CURRENT_TOTAL,
        "as_of": "2026-06-01",
        "note_ko": CURRENT_NOTE,
        "vacancies": 0,
    }


def build() -> Dict[str, Any]:
    path = ensure_xlsx()
    wb = load_workbook(path, read_only=True, data_only=True)
    smd = parse_smd(wb)
    heads = parse_pr_heads(wb)
    wb.close()
    used = set(CURRENT_BY_PARTY)
    by_party = Counter()
    for d in smd:
        for c in d["candidates"]:
            used.add(c["party_abbr"])
            by_party[c["party_abbr"]] += 1
    used.update(h["party_abbr"] for h in heads)
    print("SMD by party:", dict(by_party), "sum", sum(by_party.values()))
    doc = {
        "schema": "elections_contest_v1",
        "event_id": EVENT_ID,
        "iso3": "DEU",
        "date": "2026-09-20",
        "system_id": "deu-landtag-mv",
        "as_of": AS_OF,
        "sources": [XLSX, PDF, LAIV, PLENUM, FRACTIONS, LANDTAG_HOME, LANDTAG_WAHL, WIKI_DE],
        "parties": parties_block(used),
        "columns": [
            {
                "key": "smd",
                "label_ko": "소선거구",
                "current": current_block(),
                "contested_ko": "36개 지역구 전원. 투표 2026-09-20, 집계 전.",
                "seat_note_ko": "연동형. 지역구 당선은 명부 몫에서 뺀다. 후보는 Wahlheft 4/2026 등록자 263명(정당 지역구 258 + 무소속 5).",
                "districts": smd,
            },
            {
                "key": "pr",
                "label_ko": "비례대표",
                "current": current_block(),
                "contested_ko": "법정 최소 35석 명부. 초과의석·보정의석으로 늘어날 수 있다.",
                "seat_note_ko": "연동형. 5% 저지선. Hare/Niemeyer. 기본의석 조항 없음.",
                "districts": [pr_district(heads)],
            },
        ],
    }
    return compact(doc)


def upsert_index() -> None:
    path = PUBLIC / "elections_contests_index_v1.json"
    if path.exists():
        idx = json.loads(path.read_text(encoding="utf-8"))
    else:
        idx = {
            "schema": "elections_contests_index_v1",
            "null_policy": {"없음": "해당 없음", "불명": "미확정"},
            "events": [],
        }
    idx["as_of"] = AS_OF
    idx.setdefault("null_policy", {"없음": "해당 없음", "불명": "미확정"})
    row = {
        "event_id": EVENT_ID,
        "iso3": "DEU",
        "date": "2026-09-20",
        "path": f"elections_contests/{EVENT_ID}.json",
        "status": "complete",
        "as_of": AS_OF,
    }
    events = [e for e in idx.get("events", []) if e.get("event_id") != EVENT_ID]
    events.append(row)
    events.sort(key=lambda e: (e.get("date") or "", e.get("event_id") or ""))
    idx["events"] = events
    path.write_text(json.dumps(idx, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    doc = build()
    CONTESTS_DIR.mkdir(parents=True, exist_ok=True)
    smd = doc["columns"][0].pop("districts")
    smd_rel = f"elections_contests/{EVENT_ID}.smd.districts.json"
    payload = json.dumps(
        {"schema": "elections_contest_districts_v1", "districts": smd},
        ensure_ascii=False,
        indent=2,
    )
    inline = json.dumps(
        {**doc, "columns": [{**doc["columns"][0], "districts": smd}, doc["columns"][1]]},
        ensure_ascii=False,
    )
    if len(inline.encode("utf-8")) > 1_000_000:
        (PUBLIC / smd_rel).write_text(payload + "\n", encoding="utf-8")
        doc["columns"][0]["districts_path"] = smd_rel
        out_districts = smd
    else:
        doc["columns"][0]["districts"] = smd
        out_districts = smd
    out = CONTESTS_DIR / f"{EVENT_ID}.json"
    out.write_text(json.dumps(doc, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    upsert_index()
    pr = doc["columns"][1]["districts"]
    print(
        f"wrote {out} ({out.stat().st_size} bytes); "
        f"SMD {len(out_districts)} / {sum(len(d['candidates']) for d in out_districts)}; "
        f"PR {len(pr)} / {sum(len(d['candidates']) for d in pr)}"
    )


if __name__ == "__main__":
    main()
