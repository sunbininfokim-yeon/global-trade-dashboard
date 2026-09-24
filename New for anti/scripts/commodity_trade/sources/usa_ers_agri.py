"""US agri export volumes from ERS FATUS top-10 markets workbook.

Uses explicit **World total** rows (metric tons) per commodity per month sheet.
This is U.S. export total — NOT global production/trade of all countries.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any
from urllib.request import Request, urlopen

from openpyxl import load_workbook

URL = (
    "https://www.ers.usda.gov/media/5031/"
    "top-10-us-agricultural-export-markets-for-wheat-corn-soybeans-and-cotton-by-volume.xlsx"
)

COMMODITY_MAP = {
    "Soybeans": "soybeans",
    "Corn": "corn",
    "Wheat, unmilled": "wheat",
    "Wheat": "wheat",
    "Cotton, excluding linters": "cotton",
    "Cotton": "cotton",
}

MONTH_NAME = {
    "jan": "01", "january": "01",
    "feb": "02", "february": "02",
    "mar": "03", "march": "03",
    "apr": "04", "april": "04",
    "may": "05",
    "jun": "06", "june": "06",
    "jul": "07", "july": "07",
    "aug": "08", "august": "08",
    "sep": "09", "september": "09",
    "oct": "10", "october": "10",
    "nov": "11", "november": "11",
    "dec": "12", "december": "12",
}


def fetch(cache_dir: Path) -> Path:
    cache_dir.mkdir(parents=True, exist_ok=True)
    dest = cache_dir / "ERS_top10_export_markets_volume.xlsx"
    if dest.exists() and dest.stat().st_size > 10000:
        return dest
    data = urlopen(Request(URL, headers={"User-Agent": "commodity-trade-ers/1.0"}), timeout=90).read()
    dest.write_bytes(data)
    return dest


def _sheet_month_year(ws) -> tuple[str | None, int | None]:
    """Return (MM, year_for_current_month_col) from header rows."""
    r2 = [c.value for c in ws[2]]
    r3 = [c.value for c in ws[3]]
    # monthly label often r2[3] e.g. 'June'
    label = str(r2[3] or r2[2] or "").strip().lower()
    mm = None
    for k, v in MONTH_NAME.items():
        if label.startswith(k) or k in label:
            mm = v
            break
    # current year for monthly column (col index 4 -> r3[4] if 1-based)
    # openpyxl row values: index 0 = col A
    year = None
    for idx in (4, 3, 2):
        if idx < len(r3) and isinstance(r3[idx], (int, float)):
            year = int(r3[idx])
            break
    # sheet title fallback: 'Aug. update of Jun. data'
    return mm, year


def _parse_sheet(ws) -> list[dict]:
    mm, year = _sheet_month_year(ws)
    if not mm or not year:
        # try from sheet title via parent
        return []
    month = f"{year}-{mm}"
    points = []
    current_cid = None
    for row in ws.iter_rows(min_row=5, values_only=True):
        name = row[0]
        if name is None:
            continue
        name_s = str(name).strip()
        if name_s in COMMODITY_MAP:
            current_cid = COMMODITY_MAP[name_s]
            continue
        if name_s.lower().startswith("world total") and current_cid:
            # monthly current year column = index 4
            val = row[4] if len(row) > 4 else None
            if val is None:
                val = row[3] if len(row) > 3 else None
            try:
                v = float(val)
            except (TypeError, ValueError):
                current_cid = None
                continue
            points.append(
                {
                    "month": month,
                    "value": v,
                    "unit": "metric_tons",
                    "product_code": current_cid,
                    "source": "ers_fatus_top10_world_total",
                    "hs_bridge": {
                        "wheat": "1001",
                        "corn": "1005",
                        "soybeans": "1201",
                        "cotton": "5201",
                    }.get(current_cid),
                    "hs_relation": "alias",
                    "code_system": "ERS_FATUS",
                    "note": "USA export World total (not global multi-country sum)",
                }
            )
            current_cid = None
    return points


def load_us_agri_exports(*, cache_dir: Path) -> dict[str, Any]:
    try:
        path = fetch(cache_dir)
    except Exception as exc:
        return {"available": False, "reason": str(exc), "series": {}}

    wb = load_workbook(path, data_only=True)
    by_cid: dict[str, dict[str, dict]] = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        # improve month from sheet name if header fails
        pts = _parse_sheet(ws)
        if not pts:
            # fallback month from sheet name
            m = re.search(
                r"of\s+(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)",
                sheet_name,
                re.I,
            )
            if not m:
                continue
            mm = MONTH_NAME[m.group(1).lower()[:3]]
            # year from row3 col E
            r3 = [c.value for c in ws[3]]
            year = None
            for idx in (4, 3, 2):
                if idx < len(r3) and isinstance(r3[idx], (int, float)):
                    year = int(r3[idx])
                    break
            if not year:
                continue
            # re-parse with forced month by mutating logic inline
            current_cid = None
            month = f"{year}-{mm}"
            for row in ws.iter_rows(min_row=5, values_only=True):
                name = row[0]
                if name is None:
                    continue
                name_s = str(name).strip()
                if name_s in COMMODITY_MAP:
                    current_cid = COMMODITY_MAP[name_s]
                    continue
                if name_s.lower().startswith("world total") and current_cid:
                    val = row[4] if len(row) > 4 else row[3]
                    try:
                        v = float(val)
                    except (TypeError, ValueError):
                        current_cid = None
                        continue
                    pts.append(
                        {
                            "month": month,
                            "value": v,
                            "unit": "metric_tons",
                            "product_code": current_cid,
                            "source": "ers_fatus_top10_world_total",
                            "hs_bridge": {
                                "wheat": "1001",
                                "corn": "1005",
                                "soybeans": "1201",
                                "cotton": "5201",
                            }.get(current_cid),
                            "hs_relation": "alias",
                            "code_system": "ERS_FATUS",
                            "note": "USA export World total (not global multi-country sum)",
                        }
                    )
                    current_cid = None
        for p in pts:
            cid = p["product_code"]
            by_cid.setdefault(cid, {})[p["month"]] = p

    series = {
        cid: {"USA": [by_cid[cid][m] for m in sorted(by_cid[cid].keys())]}
        for cid in by_cid
    }
    return {
        "available": bool(series),
        "source": "ers_fatus",
        "source_url": URL,
        "series": series,
        "note_ko": (
            "미국 수출 World total(톤). 전 세계 합산이 아니라 '미국→세계' 수출 총량. "
            "브라질 kg와 단위 다름 — 합산 금지."
        ),
        "world_semantics": "reporter_export_to_world",
    }
