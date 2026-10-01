"""Australia ABS International Trade in Goods — Metal Ore Mining export FOB.

Table 32a ANZSIC industry series (monthly AUD millions).
Iron ore dominates metal ore exports; labeled as proxy not pure HS 2601.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.request import Request, urlopen

from openpyxl import load_workbook

# Jun-2026 release path (update when ABS rolls path)
DEFAULT_URL = (
    "https://www.abs.gov.au/statistics/economy/international-trade/"
    "international-trade-goods/jun-2026/5368032a.xlsx"
)
SERIES_ID = "A3343654T"  # Metal Ore Mining ;


def fetch_workbook(cache_dir: Path, url: str = DEFAULT_URL) -> Path:
    cache_dir.mkdir(parents=True, exist_ok=True)
    dest = cache_dir / "ABS_5368032a.xlsx"
    if dest.exists() and dest.stat().st_size > 10000:
        return dest
    data = urlopen(Request(url, headers={"User-Agent": "commodity-trade-abs/1.0"}), timeout=90).read()
    dest.write_bytes(data)
    return dest


def load_iron_ore_proxy(*, cache_dir: Path) -> dict[str, Any]:
    try:
        path = fetch_workbook(cache_dir)
    except Exception as exc:
        return {"available": False, "reason": str(exc), "series": {}}

    wb = load_workbook(path, data_only=True)
    ws = wb["Data1"]
    # row 10 = Series ID
    headers = [c.value for c in ws[10]]
    try:
        col = headers.index(SERIES_ID)
    except ValueError:
        return {"available": False, "reason": f"series {SERIES_ID} not found", "series": {}}

    pts = []
    for row in ws.iter_rows(min_row=11, values_only=True):
        t = row[0]
        v = row[col]
        if not isinstance(t, datetime) or v is None:
            continue
        try:
            val = float(v)
        except (TypeError, ValueError):
            continue
        pts.append(
            {
                "month": f"{t.year:04d}-{t.month:02d}",
                "value": val,
                "unit": "AUD_million_fob",
                "product_code": "ANZSIC_Metal_Ore_Mining",
                "source": "abs_5368032a",
                "hs_bridge": "2601",
                "hs_relation": "industry_proxy_not_hs_pure",
                "code_system": "ANZSIC",
                "note": "Metal ore mining FOB; iron ore dominant but not HS2601-only",
            }
        )
    pts.sort(key=lambda p: p["month"])
    return {
        "available": bool(pts),
        "source": "abs_merchandise_trade",
        "source_url": DEFAULT_URL,
        "series": {"iron_ore": {"AUS": pts}},
        "note_ko": (
            "호주 ABS 표32a Metal Ore Mining 수출 FOB(백만 AUD). "
            "철광 비중이 크나 HS2601 순수 톤이 아님 — proxy. 단위 혼용 주의."
        ),
    }
