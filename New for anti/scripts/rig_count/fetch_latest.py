#!/usr/bin/env python3
"""Refresh NAM_latest.xlsx / WW_latest.xlsx from rigcount.bakerhughes.com.

Baker Hughes has no API. Each report page (na-rig-count, intl-rig-count)
links to the current workbook under a stable link TEXT ("... - New Report")
but an opaque /static-files/<uuid> URL that changes silently whenever a new
report is published -- so this scrapes the link off the page rather than
guessing a URL pattern. The download itself 403s from a bare requests.get /
curl (Akamai blocks the default client fingerprint); a normal browser
User-Agent plus a Referer pointing back at the report page is enough to pass,
verified by hand against the live site on 2026-08-21 (see build_rig_count_v1.py
for the wider pipeline this feeds).

Nothing here touches raw/ unless a download both succeeds AND opens as a
valid workbook with the expected sheet -- a scrape that comes back empty (site
markup changed) or a corrupt download leaves the previous good file in place
and exits non-zero, rather than silently breaking the next build.

Usage:
    python fetch_latest.py                 # refresh both, overwrite raw/
    python fetch_latest.py --dry-run        # report what would change, no writes
    python fetch_latest.py --print-stats
"""
from __future__ import annotations

import argparse
import hashlib
import re
import sys
import tempfile
import urllib.error
import urllib.request
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover - dependency hint
    print("openpyxl required:  pip install openpyxl", file=sys.stderr)
    raise

RAW = Path(__file__).resolve().parent / "raw"

UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36")

# (report page, link-text substring that identifies the current workbook link
#  among the several download links each page carries, target raw/ filename,
#  sheet the built JSON expects to find inside it).
TARGETS = [
    ("https://rigcount.bakerhughes.com/na-rig-count", "North America Rig Count Report - New Report",
     "NAM_latest.xlsx", "NAM Monthly"),
    ("https://rigcount.bakerhughes.com/intl-rig-count", "Worldwide Rig Count Report - New Report",
     "WW_latest.xlsx", "WW Monthly"),
]

LINK_RE = re.compile(
    r'<a\b[^>]*\bhref="(?P<href>/static-files/[0-9a-fA-F-]+)"[^>]*>(?P<text>[^<]*)</a>',
    re.IGNORECASE,
)


def _get(url: str, referer: str | None = None) -> bytes:
    headers = {"User-Agent": UA, "Accept": "*/*"}
    if referer:
        headers["Referer"] = referer
    req = urllib.request.Request(url, headers=headers)
    with urllib.request.urlopen(req, timeout=30) as resp:
        return resp.read()


def find_report_url(page_url: str, link_text: str) -> str | None:
    """Scrape the current download URL for a report off its listing page.

    Matches by link TEXT, not position -- the page lists several downloads
    (older history files, FAQs, methodology notices) and only the exact
    label Baker Hughes uses for "this week/month's report" identifies the
    one that changes on a schedule.
    """
    html = _get(page_url).decode("utf-8", errors="replace")
    for m in LINK_RE.finditer(html):
        if link_text in m.group("text"):
            href = m.group("href")
            return href if href.startswith("http") else f"https://rigcount.bakerhughes.com{href}"
    return None


def refresh_one(page_url: str, link_text: str, filename: str, sheet: str, dry_run: bool) -> dict:
    result = {"filename": filename, "status": "unchanged"}
    report_url = find_report_url(page_url, link_text)
    if not report_url:
        result["status"] = "error"
        result["error"] = f"no link containing {link_text!r} found on {page_url}"
        return result
    result["source_url"] = report_url

    try:
        data = _get(report_url, referer=page_url)
    except urllib.error.HTTPError as err:
        result["status"] = "error"
        result["error"] = f"download failed: HTTP {err.code}"
        return result

    if len(data) < 10_000:
        result["status"] = "error"
        result["error"] = f"download suspiciously small ({len(data)} bytes), not writing"
        return result

    # Verify before touching raw/: a workbook that opens and has the sheet
    # this pipeline reads from is the bar, not just "some bytes arrived".
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp.write(data)
        tmp_path = Path(tmp.name)
    try:
        wb = openpyxl.load_workbook(tmp_path, read_only=True, data_only=True)
        if sheet not in wb.sheetnames:
            result["status"] = "error"
            result["error"] = f"downloaded workbook has no {sheet!r} sheet (sheets: {wb.sheetnames})"
            tmp_path.unlink(missing_ok=True)
            return result
        wb.close()
    except Exception as err:  # noqa: BLE001 - any openpyxl failure means "not a valid workbook"
        result["status"] = "error"
        result["error"] = f"downloaded file will not open as .xlsx: {err}"
        tmp_path.unlink(missing_ok=True)
        return result

    target = RAW / filename
    new_hash = hashlib.sha256(data).hexdigest()
    old_hash = hashlib.sha256(target.read_bytes()).hexdigest() if target.exists() else None
    if old_hash == new_hash:
        result["status"] = "unchanged"
        tmp_path.unlink(missing_ok=True)
        return result

    result["status"] = "updated"
    result["bytes"] = len(data)
    result["previous_bytes"] = target.stat().st_size if target.exists() else None
    if dry_run:
        tmp_path.unlink(missing_ok=True)
    else:
        tmp_path.replace(target)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--dry-run", action="store_true", help="scrape and verify only, do not write raw/")
    parser.add_argument("--print-stats", action="store_true")
    args = parser.parse_args()

    had_error = False
    for page_url, link_text, filename, sheet in TARGETS:
        r = refresh_one(page_url, link_text, filename, sheet, args.dry_run)
        if r["status"] == "error":
            had_error = True
            print(f"[fetch_latest] {filename}: ERROR - {r['error']}", file=sys.stderr)
        elif args.print_stats or r["status"] == "updated":
            extra = f" ({r.get('previous_bytes')} -> {r.get('bytes')} bytes)" if r["status"] == "updated" else ""
            print(f"[fetch_latest] {filename}: {r['status']}{extra}"
                  f"{'  [dry-run, not written]' if args.dry_run and r['status'] == 'updated' else ''}")

    return 1 if had_error else 0


if __name__ == "__main__":
    raise SystemExit(main())
